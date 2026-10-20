from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.db.models import ProtectedError, Q
from django.forms import modelform_factory
from django.core.paginator import Paginator
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_POST
from io import BytesIO
from openpyxl import Workbook

from core.decorators import require
from core.delegation_scope import can_select_delegation, scope_queryset

from .models import Delegation, Functionary, FunctionaryRole, Position, Role

# ---------- Delegation ----------
DelegationForm = modelform_factory(Delegation, fields=["name", "state", "field"])


@login_required
@require_POST
def select_delegation(request):
    if not can_select_delegation(request.user):
        raise PermissionDenied

    delegation_id = request.POST.get("delegation")
    if delegation_id:
        delegation = get_object_or_404(Delegation, pk=delegation_id)
        request.session["selected_delegation_id"] = delegation.pk
    else:
        request.session.pop("selected_delegation_id", None)

    next_url = request.POST.get("next")
    if not url_has_allowed_host_and_scheme(
        next_url, allowed_hosts={request.get_host()}
    ):
        return redirect("dashboard_redirect")
    return redirect(next_url)


@require("personal.view_delegation")
def delegation_list(request):
    qs = _filtered_delegations(request)
    if not request.GET.get("all") and not request.GET.get("state"):
        qs = qs.filter(state="ACTIVO")
    page_obj = Paginator(qs, 10).get_page(request.GET.get("page"))
    query_params = request.GET.copy()
    query_params.pop("page", None)
    export_params = query_params.copy()
    all_params = query_params.copy()
    all_params["all"] = "1"
    all_params.pop("state", None)
    active_params = query_params.copy()
    active_params.pop("all", None)
    active_params["state"] = "ACTIVO"
    return render(request, "personal/delegation_list.html", {
        "object_list": page_obj.object_list,
        "page_obj": page_obj,
        "pagination_query": query_params.urlencode(),
        "export_query": export_params.urlencode(),
        "all_query": all_params.urlencode(),
        "active_query": active_params.urlencode(),
        "showing_all": bool(request.GET.get("all")),
        "states": Delegation.STATES_CHOICES,
        "fields": Delegation.FIELD_CHOICES,
    })


@require("personal.view_delegation")
def delegation_export(request):
    delegations = _filtered_delegations(request)
    if not request.GET.get("all") and not request.GET.get("state"):
        delegations = delegations.filter(state="ACTIVO")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Delegaciones"
    worksheet.append(["Nombre", "Estado", "Sector"])
    for delegation in delegations:
        worksheet.append([
            _excel_safe_text(delegation.name),
            delegation.get_state_display(),
            delegation.get_field_display(),
        ])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.column_dimensions["A"].width = 36
    worksheet.column_dimensions["B"].width = 18
    worksheet.column_dimensions["C"].width = 18

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="delegaciones.xlsx"'
    return response


def _filtered_delegations(request):
    qs = scope_queryset(request, Delegation.objects.all(), lookup="pk")
    if search := request.GET.get("q", "").strip():
        qs = qs.filter(name__icontains=search)
    if state := request.GET.get("state", "").strip():
        qs = qs.filter(state=state)
    if field := request.GET.get("field", "").strip():
        qs = qs.filter(field=field)
    return qs.order_by("name", "pk")


@require("personal.add_delegation")
def delegation_create(request):
    form = DelegationForm(request.POST or None)
    if form.is_valid():
        form.save()
        messages.success(request, "Registro creado correctamente.")
        return redirect("personal:delegation_list")
    return render(request, "personal/delegation_form.html", {"form": form})


@require("personal.change_delegation")
def delegation_update(request, pk):
    obj = get_object_or_404(
        scope_queryset(request, Delegation.objects.all(), lookup="pk"), pk=pk
    )
    form = DelegationForm(request.POST or None, instance=obj)
    if form.is_valid():
        form.save()
        messages.success(request, "Registro actualizado correctamente.")
        return redirect("personal:delegation_list")
    return render(request, "personal/delegation_form.html", {"form": form, "object": obj})


@require("personal.change_delegation")
def delegation_delete(request, pk):
    """No borra: marca el registro como inactivo."""
    obj = get_object_or_404(
        scope_queryset(request, Delegation.objects.all(), lookup="pk"), pk=pk
    )
    if request.method == "POST":
        obj.state = 'INACTIVO'
        obj.save(update_fields=["state"])
        messages.success(request, "Registro desactivado.")
        return redirect("personal:delegation_list")
    return render(request, "personal/delegation_confirm_delete.html", {"object": obj})

# ---------- Position ----------
PositionForm = modelform_factory(Position, fields=["name", "validity"])


@require("personal.view_position")
def position_list(request):
    qs = _filtered_positions(request)
    if not request.GET.get("all"):
        qs = qs.filter(validity=True)
    page_obj = Paginator(qs, 10).get_page(request.GET.get("page"))
    query_params = request.GET.copy()
    query_params.pop("page", None)
    return render(request, "personal/position_list.html", {
        "object_list": page_obj.object_list,
        "page_obj": page_obj,
        "pagination_query": query_params.urlencode(),
        "showing_all": bool(request.GET.get("all")),
    })


@require("personal.view_position")
def position_export(request):
    positions = _filtered_positions(request)
    if not request.GET.get("all"):
        positions = positions.filter(validity=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Cargos"
    worksheet.append(["Cargo", "Estado"])
    for position in positions:
        worksheet.append([
            _excel_safe_text(position.name),
            "Activo" if position.validity else "Inactivo",
        ])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.column_dimensions["A"].width = 40
    worksheet.column_dimensions["B"].width = 18

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="cargos.xlsx"'
    return response


def _filtered_positions(request):
    qs = Position.objects.all().order_by("name", "pk")
    if search := request.GET.get("q", "").strip():
        qs = qs.filter(name__icontains=search)
    return qs


@require("personal.add_position")
def position_create(request):
    form = PositionForm(request.POST or None)
    if form.is_valid():
        form.save()
        messages.success(request, "Cargo creado correctamente.")
        return redirect("personal:position_list")
    return render(request, "personal/position_form.html", {"form": form})


@require("personal.change_position")
def position_update(request, pk):
    obj = get_object_or_404(Position, pk=pk)
    form = PositionForm(request.POST or None, instance=obj)
    if form.is_valid():
        form.save()
        messages.success(request, "Cargo actualizado correctamente.")
        return redirect("personal:position_list")
    return render(request, "personal/position_form.html", {"form": form, "object": obj})


@require("personal.change_position")
def position_delete(request, pk):
    """Desactiva el cargo sin eliminar sus asociaciones con funcionarios."""
    obj = get_object_or_404(Position, pk=pk)
    if request.method == "POST":
        obj.validity = False
        obj.save(update_fields=["validity"])
        messages.success(request, "Cargo desactivado.")
        return redirect("personal:position_list")
    return render(request, "personal/position_confirm_delete.html", {"object": obj})

# ---------- Functionary ----------
FunctionaryForm = modelform_factory(Functionary, fields=["names", "lastnames", "delegation", "position"])


@require("personal.view_functionary")
def functionary_list(request):
    qs = _filtered_functionaries(request)
    page_obj = Paginator(qs, 10).get_page(request.GET.get("page"))
    query_params = request.GET.copy()
    query_params.pop("page", None)
    return render(request, "personal/functionary_list.html", {
        "object_list": page_obj.object_list,
        "page_obj": page_obj,
        "pagination_query": query_params.urlencode(),
        "showing_all": bool(request.GET.get("all")),
        "delegations": scope_queryset(
            request, Delegation.objects.filter(state="ACTIVO"), lookup="pk"
        ),
        "positions": Position.objects.filter(validity=True),
    })


@require("personal.view_functionary")
def functionary_export(request):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Funcionarios"
    worksheet.append(["Nombres", "Apellidos", "Delegación", "Cargo"])
    for functionary in _filtered_functionaries(request):
        worksheet.append([
            _excel_safe_text(functionary.names),
            _excel_safe_text(functionary.lastnames),
            _excel_safe_text(functionary.delegation.name),
            _excel_safe_text(functionary.position.name),
        ])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column, width in zip(("A", "B", "C", "D"), (24, 24, 28, 28)):
        worksheet.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="funcionarios.xlsx"'
    return response


def _filtered_functionaries(request):
    qs = scope_queryset(
        request,
        Functionary.objects.select_related("delegation", "position"),
        lookup="delegation_id",
    )
    if search := request.GET.get("q", "").strip():
        qs = qs.filter(
            Q(names__icontains=search)
            | Q(lastnames__icontains=search)
            | Q(delegation__name__icontains=search)
            | Q(position__name__icontains=search)
        )
    if delegation := request.GET.get("delegation"):
        qs = qs.filter(delegation_id=delegation)
    if position := request.GET.get("position"):
        qs = qs.filter(position_id=position)
    return qs.order_by("lastnames", "names", "pk")


def _excel_safe_text(value):
    if value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


@require("personal.add_functionary")
def functionary_create(request):
    form = FunctionaryForm(request.POST or None)
    form.fields["delegation"].queryset = scope_queryset(
        request, Delegation.objects.all(), lookup="pk"
    )
    if form.is_valid():
        form.save()
        messages.success(request, "Registro creado correctamente.")
        return redirect("personal:functionary_list")
    return render(request, "personal/functionary_form.html", {"form": form})


@require("personal.change_functionary")
def functionary_update(request, pk):
    obj = get_object_or_404(_scoped_functionaries(request), pk=pk)
    form = FunctionaryForm(request.POST or None, instance=obj)
    form.fields["delegation"].queryset = scope_queryset(
        request, Delegation.objects.all(), lookup="pk"
    )
    if form.is_valid():
        form.save()
        messages.success(request, "Registro actualizado correctamente.")
        return redirect("personal:functionary_list")
    return render(request, "personal/functionary_form.html", {"form": form, "object": obj})


@require("personal.delete_functionary")
def functionary_delete(request, pk):
    obj = get_object_or_404(_scoped_functionaries(request), pk=pk)
    if request.method == "POST":
        try:
            obj.delete()
            messages.success(request, "Registro eliminado.")
        except ProtectedError:
            messages.error(request, "No se puede eliminar: tiene registros asociados.")
        return redirect("personal:functionary_list")
    return render(request, "personal/functionary_confirm_delete.html", {"object": obj})


def _scoped_functionaries(request):
    return scope_queryset(request, Functionary.objects.all(), lookup="delegation_id")


# ---------- Role ----------
RoleForm = modelform_factory(Role, fields=["role_name"])


@require("personal.view_role")
def role_list(request):
    return render(request, "personal/role_list.html", {"object_list": Role.objects.all()})


@require("personal.add_role")
def role_create(request):
    form = RoleForm(request.POST or None)
    if form.is_valid():
        form.save()
        messages.success(request, "Registro creado correctamente.")
        return redirect("personal:role_list")
    return render(request, "personal/role_form.html", {"form": form})


@require("personal.change_role")
def role_update(request, pk):
    obj = get_object_or_404(Role, pk=pk)
    form = RoleForm(request.POST or None, instance=obj)
    if form.is_valid():
        form.save()
        messages.success(request, "Registro actualizado correctamente.")
        return redirect("personal:role_list")
    return render(request, "personal/role_form.html", {"form": form, "object": obj})


@require("personal.delete_role")
def role_delete(request, pk):
    obj = get_object_or_404(Role, pk=pk)
    if request.method == "POST":
        # FunctionaryRole usa CASCADE: sin este chequeo se perderían las asignaciones en silencio.
        if obj.functionaryrole_set.exists():
            messages.error(request, "No se puede eliminar: hay funcionarios con este rol.")
        else:
            obj.delete()
            messages.success(request, "Registro eliminado.")
        return redirect("personal:role_list")
    return render(request, "personal/role_confirm_delete.html", {"object": obj})


# ---------- FunctionaryRole (asignar roles a un funcionario) ----------
FunctionaryRoleForm = modelform_factory(FunctionaryRole, fields=["functionary", "role"])


@require("personal.view_functionaryrole")
def functionary_role_list(request):
    qs = _filtered_functionary_roles(request)
    page_obj = Paginator(qs, 10).get_page(request.GET.get("page"))
    query_params = request.GET.copy()
    query_params.pop("page", None)
    return render(request, "personal/functionary_role_list.html", {
        "object_list": page_obj.object_list,
        "page_obj": page_obj,
        "pagination_query": query_params.urlencode(),
        "functionaries": _scoped_functionaries(request),
        "roles": Role.objects.order_by("role_name"),
    })


@require("personal.view_functionaryrole")
def functionary_role_export(request):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Roles de funcionarios"
    worksheet.append(["Nombres", "Apellidos", "Delegación", "Cargo", "Rol"])
    for assignment in _filtered_functionary_roles(request):
        worksheet.append([
            _excel_safe_text(assignment.functionary.names),
            _excel_safe_text(assignment.functionary.lastnames),
            _excel_safe_text(assignment.functionary.delegation.name),
            _excel_safe_text(assignment.functionary.position.name),
            _excel_safe_text(assignment.role.role_name),
        ])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column, width in zip(("A", "B", "C", "D", "E"), (22, 24, 28, 28, 24)):
        worksheet.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="roles_funcionarios.xlsx"'
    return response


def _filtered_functionary_roles(request):
    qs = scope_queryset(
        request,
        FunctionaryRole.objects.select_related(
            "functionary__delegation",
            "functionary__position",
            "role",
        ),
    )
    if search := request.GET.get("q", "").strip():
        qs = qs.filter(
            Q(functionary__names__icontains=search)
            | Q(functionary__lastnames__icontains=search)
            | Q(functionary__delegation__name__icontains=search)
            | Q(functionary__position__name__icontains=search)
            | Q(role__role_name__icontains=search)
        )
    if functionary := request.GET.get("functionary"):
        qs = qs.filter(functionary_id=functionary)
    if role := request.GET.get("role"):
        qs = qs.filter(role_id=role)
    return qs.order_by(
        "functionary__lastnames",
        "functionary__names",
        "role__role_name",
        "pk",
    )


@require("personal.add_functionaryrole")
def functionary_role_create(request):
    # Permite abrir el formulario con el funcionario ya elegido: ?functionary=3
    initial = {"functionary": request.GET.get("functionary")}
    form = FunctionaryRoleForm(request.POST or None, initial=initial)
    form.fields["functionary"].queryset = _scoped_functionaries(request)
    if form.is_valid():
        form.save()
        messages.success(request, "Rol asignado correctamente.")
        return redirect("personal:functionary_role_list")
    return render(request, "personal/functionary_role_form.html", {"form": form})


# No hay "editar": para cambiar un rol se quita la asignación y se crea otra.
@require("personal.delete_functionaryrole")
def functionary_role_delete(request, pk):
    obj = get_object_or_404(
        scope_queryset(
            request,
            FunctionaryRole.objects.select_related("functionary"),
        ),
        pk=pk,
    )
    if request.method == "POST":
        obj.delete()
        messages.success(request, "Rol quitado.")
        return redirect("personal:functionary_role_list")
    return render(request, "personal/functionary_role_confirm_delete.html", {"object": obj})