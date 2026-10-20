from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .models import Delegation, Functionary, FunctionaryRole, Position, Role


class FunctionaryListTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        user_model = get_user_model()
        cls.admin = user_model.objects.create_superuser(
            username="functionary-admin",
            email="admin@example.com",
            password="test-password",
        )
        cls.delegation = Delegation.objects.create(
            name="Central",
            state="ACTIVO",
            field="URBANO",
        )
        cls.position = Position.objects.create(name="Inspector", validity=True)
        Functionary.objects.bulk_create([
            Functionary(
                names=f"Nombre {number}",
                lastnames=f"Apellido {number:02}",
                delegation=cls.delegation,
                position=cls.position,
            )
            for number in range(12)
        ])

    def setUp(self):
        self.client.force_login(self.admin)

    def test_functionary_list_is_paginated(self):
        response = self.client.get(reverse("personal:functionary_list"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["object_list"]), 10)
        self.assertEqual(response.context["page_obj"].paginator.count, 12)
        self.assertContains(response, "data-modal-form")

    def test_pagination_preserves_filters(self):
        response = self.client.get(
            reverse("personal:functionary_list"),
            {"delegation": self.delegation.pk, "page": 2},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["object_list"]), 2)
        self.assertEqual(
            response.context["pagination_query"],
            f"delegation={self.delegation.pk}",
        )

    def test_search_filters_functionaries_across_name_and_related_fields(self):
        response = self.client.get(
            reverse("personal:functionary_list"),
            {"q": "Apellido 03"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["object_list"]), 1)
        self.assertEqual(response.context["object_list"][0].lastnames, "Apellido 03")
        self.assertContains(response, "data-live-search-input")

    def test_export_contains_all_filtered_functionaries(self):
        response = self.client.get(
            reverse("personal:functionary_export"),
            {"delegation": self.delegation.pk},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(len(rows), 13)
        self.assertEqual(rows[0], ("Nombres", "Apellidos", "Delegación", "Cargo"))


class DelegationCrudTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_superuser(
            username="delegation-admin",
            email="delegation-admin@example.com",
            password="test-password",
        )

    def setUp(self):
        self.client.force_login(self.admin)

    def test_list_shows_active_delegations_and_can_show_all(self):
        active = Delegation.objects.create(
            name="Centro", state="ACTIVO", field="URBANO"
        )
        inactive = Delegation.objects.create(
            name="Rural Norte", state="INACTIVO", field="RURAL"
        )

        response = self.client.get(reverse("personal:delegation_list"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(list(response.context["object_list"]), [active])

        response = self.client.get(reverse("personal:delegation_list"), {"all": "1"})
        self.assertEqual(set(response.context["object_list"]), {active, inactive})

    def test_delegation_list_search_filters_and_pagination(self):
        selected = Delegation.objects.create(
            name="Norte Rural", state="INACTIVO", field="RURAL"
        )
        Delegation.objects.create(
            name="Norte Urbano", state="ACTIVO", field="URBANO"
        )

        response = self.client.get(
            reverse("personal:delegation_list"),
            {"all": "1", "q": "norte", "state": "INACTIVO", "field": "RURAL"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(list(response.context["object_list"]), [selected])
        self.assertEqual(
            response.context["pagination_query"],
            "all=1&q=norte&state=INACTIVO&field=RURAL",
        )
        self.assertContains(response, "data-live-search-input")
        self.assertContains(response, "data-live-search-filter")
        self.assertContains(response, "personal/delegations/export/")

    def test_delegation_export_respects_filters(self):
        Delegation.objects.create(
            name="Exportada Rural", state="INACTIVO", field="RURAL"
        )
        Delegation.objects.create(
            name="No incluir", state="ACTIVO", field="URBANO"
        )

        response = self.client.get(
            reverse("personal:delegation_export"),
            {"all": "1", "q": "Exportada", "field": "RURAL"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn('filename="delegaciones.xlsx"', response["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(
            rows,
            [
                ("Nombre", "Estado", "Sector"),
                ("Exportada Rural", "Inactivo", "Rural"),
            ],
        )

    def test_create_delegation(self):
        response = self.client.post(
            reverse("personal:delegation_create"),
            {"name": "Centro", "state": "ACTIVO", "field": "URBANO"},
        )

        self.assertRedirects(response, reverse("personal:delegation_list"))
        self.assertTrue(Delegation.objects.filter(name="Centro").exists())

    def test_update_delegation(self):
        delegation = Delegation.objects.create(
            name="Centro", state="ACTIVO", field="URBANO"
        )
        response = self.client.post(
            reverse("personal:delegation_update", args=[delegation.pk]),
            {"name": "Centro Histórico", "state": "ACTIVO", "field": "RURAL"},
        )

        self.assertRedirects(response, reverse("personal:delegation_list"))
        delegation.refresh_from_db()
        self.assertEqual(delegation.name, "Centro Histórico")
        self.assertEqual(delegation.field, "RURAL")

    def test_delete_route_deactivates_instead_of_removing(self):
        delegation = Delegation.objects.create(
            name="Centro", state="ACTIVO", field="URBANO"
        )
        response = self.client.post(
            reverse("personal:delegation_delete", args=[delegation.pk])
        )

        self.assertRedirects(response, reverse("personal:delegation_list"))
        delegation.refresh_from_db()
        self.assertEqual(delegation.state, "INACTIVO")


class DelegationScopeTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        user_model = get_user_model()
        cls.admin = user_model.objects.create_superuser(
            username="scope-admin",
            email="scope-admin@example.com",
            password="test-password",
        )
        cls.user = user_model.objects.create_user(
            username="scope-user",
            password="test-password",
        )
        cls.user.user_permissions.add(
            Permission.objects.get(
                content_type__app_label="personal",
                codename="view_functionary",
            )
        )
        cls.first_delegation = Delegation.objects.create(
            name="Delegación Uno", state="ACTIVO", field="URBANO"
        )
        cls.second_delegation = Delegation.objects.create(
            name="Delegación Dos", state="ACTIVO", field="RURAL"
        )
        position = Position.objects.create(name="Inspector", validity=True)
        Functionary.objects.create(
            user=cls.user,
            names="Usuario",
            lastnames="Uno",
            delegation=cls.first_delegation,
            position=position,
        )
        Functionary.objects.create(
            names="Usuario",
            lastnames="Dos",
            delegation=cls.second_delegation,
            position=position,
        )

    def test_staff_can_select_a_delegation_or_all(self):
        self.client.force_login(self.admin)
        select_url = reverse("personal:select_delegation")
        list_url = reverse("personal:functionary_list")

        response = self.client.post(
            select_url,
            {"delegation": self.second_delegation.pk, "next": list_url},
        )
        self.assertRedirects(response, list_url)
        response = self.client.get(list_url)
        self.assertEqual(len(response.context["object_list"]), 1)
        self.assertEqual(
            response.context["object_list"][0].delegation,
            self.second_delegation,
        )
        self.assertEqual(
            response.context["current_delegation"],
            self.second_delegation,
        )
        self.assertContains(response, "active-delegation")
        self.assertContains(response, "Todas las delegaciones")

        self.client.post(select_url, {"delegation": "", "next": list_url})
        response = self.client.get(list_url)
        self.assertEqual(len(response.context["object_list"]), 2)
        self.assertTrue(response.context["viewing_all_delegations"])

    def test_regular_user_is_limited_to_own_delegation(self):
        self.client.force_login(self.user)
        list_url = reverse("personal:functionary_list")
        response = self.client.get(list_url)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["object_list"]), 1)
        self.assertEqual(
            response.context["object_list"][0].delegation,
            self.first_delegation,
        )
        self.assertEqual(
            response.context["current_delegation"],
            self.first_delegation,
        )
        self.assertFalse(response.context["can_select_delegation"])
        self.assertContains(response, f"Delegación: {self.first_delegation.name}")
        self.assertNotContains(response, "active-delegation")

        response = self.client.get(
            list_url, {"delegation": self.second_delegation.pk}
        )
        self.assertFalse(
            any(
                functionary.delegation_id == self.second_delegation.pk
                for functionary in response.context["object_list"]
            )
        )

    def test_regular_user_cannot_change_global_delegation(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("personal:select_delegation"),
            {"delegation": self.second_delegation.pk},
        )

        self.assertEqual(response.status_code, 403)


class PositionCrudTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_superuser(
            username="position-admin",
            email="position-admin@example.com",
            password="test-password",
        )
        Position.objects.bulk_create([
            Position(name=f"Cargo {number:02}", validity=True)
            for number in range(12)
        ])
        cls.inactive_position = Position.objects.create(
            name="Cargo archivado", validity=False
        )

    def setUp(self):
        self.client.force_login(self.admin)

    def test_list_is_paginated_and_can_show_inactive_positions(self):
        response = self.client.get(reverse("personal:position_list"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["object_list"]), 10)
        self.assertEqual(response.context["page_obj"].paginator.count, 12)

        response = self.client.get(
            reverse("personal:position_list"), {"all": "1", "q": "archivado"}
        )
        self.assertEqual(list(response.context["object_list"]), [self.inactive_position])

    def test_create_and_update_position(self):
        response = self.client.get(reverse("personal:position_create"))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "personal/position_form.html")

        response = self.client.post(
            reverse("personal:position_create"),
            {"name": "Analista", "validity": "on"},
        )
        self.assertRedirects(response, reverse("personal:position_list"))
        position = Position.objects.get(name="Analista")
        self.assertTrue(position.validity)

        response = self.client.get(
            reverse("personal:position_update", args=[position.pk])
        )
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "personal/position_form.html")

        response = self.client.post(
            reverse("personal:position_update", args=[position.pk]),
            {"name": "Analista senior", "validity": ""},
        )
        self.assertRedirects(response, reverse("personal:position_list"))
        position.refresh_from_db()
        self.assertEqual(position.name, "Analista senior")
        self.assertFalse(position.validity)

    def test_delete_route_deactivates_position(self):
        position = Position.objects.create(name="Supervisor", validity=True)

        response = self.client.get(
            reverse("personal:position_delete", args=[position.pk])
        )
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "personal/position_confirm_delete.html")

        response = self.client.post(
            reverse("personal:position_delete", args=[position.pk])
        )

        self.assertRedirects(response, reverse("personal:position_list"))
        position.refresh_from_db()
        self.assertFalse(position.validity)
        self.assertTrue(Position.objects.filter(pk=position.pk).exists())

    def test_export_respects_filters_and_inactive_option(self):
        response = self.client.get(
            reverse("personal:position_export"),
            {"all": "1", "q": "archivado"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(rows, [("Cargo", "Estado"), ("Cargo archivado", "Inactivo")])

    def test_position_search_is_rendered_as_live_search(self):
        response = self.client.get(
            reverse("personal:position_list"),
            {"q": "archivado", "all": "1"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            [position.name for position in response.context["object_list"]],
            ["Cargo archivado"],
        )
        self.assertContains(response, "data-live-search-input")


class FunctionaryRoleTemplateTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_superuser(
            username="functionary-role-admin",
            email="functionary-role-admin@example.com",
            password="test-password",
        )
        delegation = Delegation.objects.create(
            name="Roles", state="ACTIVO", field="URBANO"
        )
        position = Position.objects.create(name="Inspector", validity=True)
        cls.functionary = Functionary.objects.create(
            names="Ana",
            lastnames="Pérez",
            delegation=delegation,
            position=position,
        )
        cls.role = Role.objects.create(role_name="Supervisor")

    def setUp(self):
        self.client.force_login(self.admin)

    def test_role_assignment_form_renders_and_preserves_prefilled_functionary(self):
        response = self.client.get(
            reverse("personal:functionary_role_create"),
            {"functionary": self.functionary.pk},
        )

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "personal/functionary_role_form.html")
        self.assertContains(response, "Asignar rol a funcionario")
        self.assertContains(
            response,
            f'<option value="{self.functionary.pk}" selected>',
            html=False,
        )
        self.assertNotContains(response, "---------")

    def test_assignment_list_and_remove_flow_render_and_work(self):
        assignment = FunctionaryRole.objects.create(
            functionary=self.functionary,
            role=self.role,
        )
        response = self.client.get(reverse("personal:functionary_role_list"))

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "personal/functionary_role_list.html")
        self.assertContains(response, "Supervisor")

        response = self.client.get(
            reverse("personal:functionary_role_delete", args=[assignment.pk])
        )
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(
            response, "personal/functionary_role_confirm_delete.html"
        )
        response = self.client.post(
            reverse("personal:functionary_role_delete", args=[assignment.pk])
        )
        self.assertRedirects(response, reverse("personal:functionary_role_list"))
        self.assertFalse(FunctionaryRole.objects.filter(pk=assignment.pk).exists())

    def test_role_assignments_support_live_search_filters_and_pagination(self):
        roles = [
            Role.objects.create(role_name=f"Rol {number:02}")
            for number in range(11)
        ]
        FunctionaryRole.objects.bulk_create([
            FunctionaryRole(functionary=self.functionary, role=role)
            for role in roles
        ])

        response = self.client.get(reverse("personal:functionary_role_list"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["object_list"]), 10)
        self.assertEqual(response.context["page_obj"].paginator.count, 11)
        self.assertContains(response, "data-live-search-input")
        self.assertContains(response, "data-live-search-filter")

        response = self.client.get(
            reverse("personal:functionary_role_list"),
            {"q": "Rol 03", "role": roles[3].pk},
        )
        self.assertEqual(len(response.context["object_list"]), 1)
        self.assertEqual(response.context["object_list"][0].role, roles[3])

    def test_role_assignment_export_respects_filters(self):
        assignment = FunctionaryRole.objects.create(
            functionary=self.functionary,
            role=self.role,
        )

        response = self.client.get(
            reverse("personal:functionary_role_export"),
            {"functionary": self.functionary.pk, "role": self.role.pk},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(
            rows,
            [
                ("Nombres", "Apellidos", "Delegación", "Cargo", "Rol"),
                (
                    assignment.functionary.names,
                    assignment.functionary.lastnames,
                    assignment.functionary.delegation.name,
                    assignment.functionary.position.name,
                    assignment.role.role_name,
                ),
            ],
        )


class RoleCatalogTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_superuser(
            username="role-catalog-admin",
            email="role-catalog-admin@example.com",
            password="role-catalog-test-password",
        )
        cls.role = Role.objects.create(role_name="Coordinador")

    def setUp(self):
        self.client.force_login(self.admin)

    def test_role_catalog_is_available_from_admin_menu(self):
        response = self.client.get(reverse("personal:role_list"))

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "personal/role_list.html")
        self.assertContains(response, "Coordinador")
        self.assertContains(response, reverse("personal:role_list"))
        self.assertContains(response, reverse("personal:functionary_role_list"))

    def test_role_create_and_update_forms_render_and_save(self):
        response = self.client.get(reverse("personal:role_create"))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "personal/role_form.html")
        self.assertContains(response, "Nombre del rol")

        response = self.client.post(
            reverse("personal:role_create"),
            {"role_name": "Revisor"},
        )
        self.assertRedirects(response, reverse("personal:role_list"))
        self.assertTrue(Role.objects.filter(role_name="Revisor").exists())

        response = self.client.post(
            reverse("personal:role_update", args=[self.role.pk]),
            {"role_name": "Coordinador general"},
        )
        self.assertRedirects(response, reverse("personal:role_list"))
        self.role.refresh_from_db()
        self.assertEqual(self.role.role_name, "Coordinador general")

    def test_role_delete_page_renders_and_unused_role_can_be_deleted(self):
        response = self.client.get(
            reverse("personal:role_delete", args=[self.role.pk])
        )
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "personal/role_confirm_delete.html")

        response = self.client.post(
            reverse("personal:role_delete", args=[self.role.pk])
        )
        self.assertRedirects(response, reverse("personal:role_list"))
        self.assertFalse(Role.objects.filter(pk=self.role.pk).exists())
