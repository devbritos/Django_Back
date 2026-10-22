from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def excel_response(filename, headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Datos'
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='AD0000')
    for row in rows:
        sheet.append([
            _safe_cell(value) for value in row
        ])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        values = [len(str(cell.value or '')) for cell in column]
        width = min(max(max(values, default=0) + 2, 12), 50)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width

    content = BytesIO()
    workbook.save(content)
    response = HttpResponse(
        content.getvalue(),
        content_type=(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        ),
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def _safe_cell(value):
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return f"'{value}"
    return value
